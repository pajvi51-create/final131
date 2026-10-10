import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = Document()

# Set margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# Title Style
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title_p.add_run('รายงานคำอธิบายซอร์สโค้ดและสถาปัตยกรรมแอปพลิเคชันอย่างละเอียด\n(Android Final Exam Guide)')
run.bold = True
run.font.size = Pt(20)
run.font.color.rgb = RGBColor(31, 78, 121)

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_sub = sub_p.add_run('วิชา: การพัฒนาแอปพลิเคชันสำหรับอุปกรณ์เคลื่อนที่ด้วย Android (Kotlin)\n')
run_sub.font.size = Pt(14)
run_sub.bold = True

info_p = doc.add_paragraph()
info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_info = info_p.add_run('ชื่อ-นามสกุล: ปาจรีย์ สุคนธชาติ  |  รหัสนักศึกษา: 67102122131\nระดับการศึกษา: ปริญญาตรี สาขาเทคโนโลยีสารสนเทศ (IT)\nGit Repository: https://github.com/pajvi51-create/final131.git\n')
run_info.font.size = Pt(12)
run_info.font.color.rgb = RGBColor(89, 89, 89)

doc.add_paragraph('='*60)

# Section 1
doc.add_heading('บทที่ 1: ภาพรวมโครงการและสถาปัตยกรรมซอฟต์แวร์', level=1)

p1 = doc.add_paragraph()
p1.add_run('โปรเจกต์นี้ได้รับการพัฒนาขึ้นเพื่อตอบโจทย์ข้อสอบไฟนอลวิชา Android Development โดยรวมการทำงานทั้ง 2 ข้อใหญ่ (30 คะแนนเต็ม) ไว้ในแอปพลิเคชันเดียวกัน (Single App) เพื่อความสะดวกลื่นไหลในการสอบนำเสนอแบบ Live Demo ต่อหน้าผู้สอน ควบคุมการสลับหน้าด้วย Bottom Navigation Bar ด้านล่าง\n')

doc.add_heading('1.1 สถาปัตยกรรม MVVM (Model-View-ViewModel)', level=2)
doc.add_paragraph('แอพพลิเคชันยึดตามหลักสถาปัตยกรรม MVVM อย่างเคร่งครัด แบ่งหน้าที่ออกเป็น 3 เลเยอร์ชัดเจน (Separation of Concerns):')

p_mvvm = doc.add_paragraph()
p_mvvm.add_run('1. Model Layer (ข้อมูล): ').bold = True
p_mvvm.add_run('ทำหน้าที่กำหนดโครงสร้างข้อมูล Data Class (ExpenseItem, Product) และ Repository Pattern ที่แยกคำสั่งเขียนไฟล์ CSV หรือคำสั่งยิง Supabase REST API ออกจากหน้าจอ UI อย่างเด็ดขาด\n')

p_mvvm.add_run('2. ViewModel Layer (ตรรกะและการคำนวณ): ').bold = True
p_mvvm.add_run('ทำหน้าที่เป็นตัวกลางคำนวณ Business Logic (คำนวณยอดสุทธิ, ส่งคำสั่ง CRUD) และถือครอง UI State โดยรันงานเบื้องหลังบน Dispatchers.IO ผ่าน viewModelScope.launch และช่วยรักษาข้อมูลไม่ให้สูญหายเมื่อหมุนหน้าจอ (Configuration Change)\n')

p_mvvm.add_run('3. View Layer (หน้าจอแสดงผล): ').bold = True
p_mvvm.add_run('พัฒนาด้วย Jetpack Compose (Material 3) นำหลัก State Hoisting มาใช้โดยแยกเป็น Stateful Composable (สื่อสารกับ ViewModel) และ Stateless Composable (วาด UI ตามสั่ง) พร้อมรองรับ Adaptive Layout 2 คอลัมน์ขนานกันในแนวนอน (Landscape Mode)\n')

# Section 2
doc.add_heading('บทที่ 2: ตารางสรุปโครงสร้างโฟลเดอร์และไฟล์ทั้งหมดในโปรเจกต์', level=1)

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False

hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'พาธไฟล์ (File Path)'
hdr_cells[1].text = 'ประเภท / Layer'
hdr_cells[2].text = 'คำอธิบายหน้าที่ความรับผิดชอบ'

for cell in hdr_cells:
    shd_xml = parse_xml(r'<w:shd {} w:fill="1F4E79"/>'.format(nsdecls("w")))
    cell._tc.get_or_add_tcPr().append(shd_xml)
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.bold = True

file_list = [
    ('data/model/ExpenseItem.kt', 'Model (ข้อ 1)', 'Data Class เก็บ Date, Time, Name, Amount, Type และฟังก์ชันแปลง CSV toCsvRow() / fromCsvRow()'),
    ('data/model/Product.kt', 'Model (ข้อ 2)', 'Data Class เก็บข้อมูลสินค้า Supabase พร้อมแอนโนเทชัน @Serializable และ @SerialName'),
    ('data/repository/ExpenseRepositoryImpl.kt', 'Repository (ข้อ 1)', 'จัดการอ่าน เขียนแบบ Append และลบไฟล์ transactions.csv บน Internal Storage (context.filesDir)'),
    ('data/repository/ProductRepositoryImpl.kt', 'Repository (ข้อ 2)', 'สื่อสารกับ Supabase REST API ครบ 4 ด้าน (Create, Read, Update +1/-1, Delete) ผ่าน HttpURLConnection'),
    ('ui/expense/ExpenseViewModel.kt', 'ViewModel (ข้อ 1)', 'คำนวณยอดสุทธิ (Net Balance = รับ - จ่าย) รันบน Dispatchers.IO และถือครอง UI State ทนต่อการหมุนจอ'),
    ('ui/expense/ExpenseTrackerScreen.kt', 'Compose View (ข้อ 1)', 'หน้าจอ UI Material 3 แสดงสรุปยอด ฟอร์มกรอกข้อมูล และ LazyColumn ประวัติ รองรับ Landscape 2 คอลัมน์'),
    ('ui/stock/ProductStockViewModel.kt', 'ViewModel (ข้อ 2)', 'จัดการงาน CRUD สต็อกสินค้าบน Supabase ควบคุม Loading/Error State รันบน Dispatchers.IO'),
    ('ui/stock/ProductStockScreen.kt', 'Compose View (ข้อ 2)', 'หน้าจอ UI Material 3 ตั้งค่า Cloud ฟอร์มเพิ่มสินค้า และ LazyColumn สินค้าพร้อมปุ่ม +1/-1 และปุ่มลบแบบ Real-time'),
    ('ui/common/InstructionScreen.kt', 'View (ศูนย์รวม)', 'หน้าคู่มือแนะนำการสอบ และปุ่มคัดลอก SQL Script สำหรับสร้างตาราง products บน Supabase SQL Editor'),
    ('MainActivity.kt', 'Main Activity', 'จุดเริ่มต้นแอป ควบคุม Bottom Navigation Bar 3 Tab สลับหน้าข้อ 1, ข้อ 2 และหน้าคู่มือ'),
    ('submission/67102122131_kotlin_java.zip', 'Deliverable Zip', 'ไฟล์บีบอัดซอร์สโค้ดภาษา Kotlin ทั้งหมดส่งอาจารย์ ตั้งชื่อตามรหัสนักศึกษา 67102122131')
]

for fpath, ftype, fdesc in file_list:
    row_cells = table.add_row().cells
    row_cells[0].text = fpath
    row_cells[1].text = ftype
    row_cells[2].text = fdesc

# Section 3
doc.add_heading('บทที่ 3: คำอธิบายซอร์สโค้ดอย่างละเอียดทีละไฟล์', level=1)

doc.add_heading('3.1 ข้อที่ 1: ระบบบันทึกรายรับ-รายจ่ายส่วนบุคคล (.CSV)', level=2)

p_exp_model = doc.add_paragraph()
p_exp_model.add_run('1. ExpenseItem.kt: ').bold = True
p_exp_model.add_run('ทำหน้าที่เป็นโครงสร้างข้อมูล Data Class เก็บวันที่ (date), เวลา (time), ชื่อรายการ (name), จำนวนเงิน (amount) และประเภทรายการ (type) มีฟังก์ชัน toCsvRow() แปลง Object เป็นข้อความ 1 บรรทัด CSV และฟังก์ชัน fromCsvRow() อ่านบรรทัด CSV แล้วตัดคำ (Split) แปลงกลับมาเป็น Object ExpenseItem เพื่อนำไปแสดงผล\n')

p_exp_repo = doc.add_paragraph()
p_exp_repo.add_run('2. ExpenseRepositoryImpl.kt: ').bold = True
p_exp_repo.add_run('ทำหน้าที่จัดการ File I/O บน Internal Storage (context.filesDir) ซึ่งเป็นพื้นที่ปลอดภัยของแอป โดยใช้ file.appendText() เพื่อบันทึกข้อมูลแบบต่อท้าย (Append) ใช้ file.useLines() อ่านไฟล์แบบ Line-by-line แปลงเป็น List<ExpenseItem> และมีฟังก์ชัน clearAllData() ลบไฟล์ transactions.csv ออกจากระบบเพื่อเคลียร์ค่า\n')

p_exp_vm = doc.add_paragraph()
p_exp_vm.add_run('3. ExpenseViewModel.kt: ').bold = True
p_exp_vm.add_run('ทำหน้าที่ถือครอง ExpenseUiState และคำนวณตรรกะทางธุรกิจ โดยคำนวณ ยอดคงเหลือสุทธิ = รายรับรวม - รายจ่ายรวม ทุกคำสั่งอ่านเขียนไฟล์ถูกรันผ่าน viewModelScope.launch บน Dispatchers.IO เพื่อไม่ให้ขัดจังหวะ Main Thread (UI ลื่นไหล ไม่กระตุก) และช่วยรักษาข้อมูลไม่ให้สูญหายเมื่อหมุนหน้าจอ\n')

p_exp_view = doc.add_paragraph()
p_exp_view.add_run('4. ExpenseTrackerScreen.kt: ').bold = True
p_exp_view.add_run('วาดหน้าจอด้วย Jetpack Compose แยกเป็น ExpenseTrackerScreen (Stateful) และ ExpenseTrackerContent (Stateless) มีการ์ด NetBalanceCard แสดงยอดสุทธิ ฟอร์ม OutlinedTextField รับชื่อรายการและจำนวนเงิน RadioButton เลือกประเภทรายการ ปุ่มบันทึก/ล้างข้อมูล และ LazyColumn แสดงประวัติย้อนหลังพร้อมสัญลักษณ์สีเขียว (+)/สีแดง (-) รองรับ Adaptive Layout 2 คอลัมน์เมื่อหมุนหน้าจอเป็นแนวนอน\n')

doc.add_heading('3.2 ข้อที่ 2: ระบบจัดการสต็อกสินค้าชุมชนบน Cloud Database (Supabase CRUD)', level=2)

p_prod_model = doc.add_paragraph()
p_prod_model.add_run('1. Product.kt: ').bold = True
p_prod_model.add_run('เป็น Data Class สินค้า (id, productName, stockQty, price) ติดตั้ง @Serializable และ @SerialName("product_name") / @SerialName("stock_qty") เพื่อให้ระบบแปลงข้อความ JSON จาก Supabase REST API เป็น Object ภาษา Kotlin อัตโนมัติ\n')

p_prod_repo = doc.add_paragraph()
p_prod_repo.add_run('2. ProductRepositoryImpl.kt: ').bold = True
p_prod_repo.add_run('สื่อสารกับ Supabase PostgreSQL Cloud Database ผ่าน HTTP REST API ด้วย HttpURLConnection ครบ 4 ด้าน:\n')
p_prod_repo.add_run('   • Read: GET /rest/v1/products?select=*&order=id.desc (เรียงจากล่าสุด)\n')
p_prod_repo.add_run('   • Create: POST /rest/v1/products (เพิ่มสินค้าใหม่)\n')
p_prod_repo.add_run('   • Update: POST ร่วมกับ X-HTTP-Method-Override: PATCH (ปรับสต็อก +1/-1)\n')
p_prod_repo.add_run('   • Delete: DELETE /rest/v1/products?id=eq.<id> (ลบสินค้าตาม ID)\n')

p_prod_vm = doc.add_paragraph()
p_prod_vm.add_run('3. ProductStockViewModel.kt: ').bold = True
p_prod_vm.add_run('ถือครอง ProductStockUiState (URL, Anon Key, ฟอร์มกรอก, รายการสินค้า, Loading/Error State) รันคำสั่งยิงคลาวด์บน Dispatchers.IO มีฟังก์ชันเพิ่มสินค้า, ปรับสต็อก +1 / -1 และลบสินค้าแบบ Real-time\n')

p_prod_view = doc.add_paragraph()
p_prod_view.add_run('4. ProductStockScreen.kt: ').bold = True
p_prod_view.add_run('วาด UI แสดงส่วนตั้งค่า Supabase Connection, ฟอร์มเพิ่มสินค้าใหม่ (Create) และ LazyColumn แสดงการ์ดสินค้าพร้อมปุ่ม +1, -1 และปุ่ม ลบ รองรับ Adaptive Layout 2 คอลัมน์ในแนวนอนเช่นกัน\n')

# Section 4
doc.add_heading('บทที่ 4: แนวทางตอบคำถามในการสอบสาธิตหน้าผู้สอน (Exam Q&A Guide)', level=1)

qa_list = [
    ('Q: ทำไมถึงเลือกทำทั้งสองข้ออยู่ในแอปพลิเคชันเดียวกัน?',
     'A: เพื่อความสะดวกลื่นไหลในการสอบนำเสนอแบบ Live Demo ต่อหน้าอาจารย์ สามารถเปิดแอปเดียวบน Emulator แล้วใช้ Bottom Navigation Bar สลับสาธิตข้อ 1 (.CSV) และข้อ 2 (Supabase) ได้ทันที โดยโครงสร้างโค้ดภายในถูกแยก Package ออกจากกันโดยเด็ดขาดตามหลัก Clean Architecture'),
    ('Q: ข้อมูลไฟล์ CSV ถูกจัดเก็บไว้ที่ไหน และปลอดภัยอย่างไร?',
     'A: จัดเก็บไว้ที่ Internal Storage (context.filesDir) ซึ่งเป็นพื้นที่จัดเก็บข้อมูลส่วนตัวของแอปพลิเคชัน แอปอื่นไม่สามารถแอบอ่านได้ และข้อมูลจะยังคงอยู่แม้ปิดแอปพลิเคชัน'),
    ('Q: ทำไมเวลาหมุนหน้าจอแล้วข้อมูลไม่สูญหาย?',
     'A: เนื่องจากใช้ ViewModel ในการถือครอง UI State ซึ่ง ViewModel มี Lifecycle ที่ยาวกว่า Activity ทำให้เมื่อเกิด Configuration Change (เช่นการหมุนจอ) Activity จะถูก Recreate ใหม่ แต่ ViewModel จะยังถูก Retain ไว้เหมือนเดิม'),
    ('Q: การใช้งาน Coroutines และ Dispatchers.IO ช่วยอะไรในแอปพลิเคชัน?',
     'A: ช่วยให้งานที่มีระยะเวลา เช่น การอ่านเขียนไฟล์ CSV หรือการยิง Network Request ไปยัง Supabase ไปทำงานอยู่บน IO Thread เบื้องหลัง ผลลัพธ์คือ Main Thread (UI Thread) จะไม่ถูกบล็อก แอปพลิเคชันทำงานได้ลื่นไหล ไม่เกิดอาการค้าง (No UI Freeze)')
]

for q, a in qa_list:
    p_q = doc.add_paragraph()
    p_q.add_run(q + '\n').bold = True
    p_q.add_run(a + '\n')

output_path = r'C:/Users/froke/AndroidStudioProjects/MyApplication6/เอกสารอธิบายซอร์สโค้ดและคู่มือการทำงาน_ปาจรีย์_67102122131.docx'
doc.save(output_path)
print('DOCX file created successfully:', output_path)
